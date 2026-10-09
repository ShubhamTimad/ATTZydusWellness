from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from modules.reports.pdf_chrome import (
    BRAND_RED,
    ChromeCanvas,
    FOOTER_BAND,
    HEADER_BAND,
    PAGE_MARGIN,
    SIG_BAND,
    XLSX_LEFT_PAD_COLS,
    apply_xlsx_data_layout,
    apply_xlsx_print_chrome,
    app_now,
    format_generated_at,
    write_xlsx_header,
    xlsx_pad_row,
)

ROW_ALT = colors.HexColor("#F8FAFC")

PDF_MARGIN = PAGE_MARGIN
PDF_BODY_FONT = 7
PDF_HEADER_FONT = 7.5
MIN_COL_WIDTH = 28
MAX_COL_WIDTH = 110

_COMPACT_HEADERS = frozenset({"plant", "status", "scans in period", "scans", "fy status", "sr. no."})
_MEDIUM_HEADERS = frozenset({"asset code", "rfid code", "rfid", "last verification", "scan time", "created at", "completed at"})
_ID_HEADERS = frozenset({
    "asset class",
    "asset category",
    "combined asset no",
    "cost center",
    "asset custodian",
    "physical condition",
    "apc-igaap",
})
_FLEXIBLE_HEADERS = frozenset({"asset name", "order number", "rejection reason", "description", "functional location", "actual location"})


def _header_key(header: str) -> str:
    return header.strip().lower()


def _column_bounds(header: str, available_width: float) -> tuple[float, float]:
    key = _header_key(header)
    if key == "asset name":
        return 64, min(available_width * 0.20, 96)
    if key == "asset class":
        return 68, 102
    if key in _FLEXIBLE_HEADERS:
        return 64, available_width * 0.32
    if key in _ID_HEADERS:
        return 56, 84
    if key in _COMPACT_HEADERS:
        return 28, 58
    if key in _MEDIUM_HEADERS or "verification" in key or "date" in key:
        return 44, 82
    return MIN_COL_WIDTH, MAX_COL_WIDTH


@dataclass(frozen=True)
class ReportDocument:
    title: str
    subtitle: str
    headers: list[str]
    rows: list[list[str]]
    summary: tuple[tuple[str, str], ...] = ()
    show_signatures: bool | None = None
    signature_blocks: tuple[tuple[str, str, str], ...] = ()
    generated_by: str = ""
    force_portrait: bool = False


def actor_display_name(db, user_id: int | None) -> str:
    if user_id is None:
        return ""
    from modules.platform.users.model import User

    user = db.get(User, user_id)
    if user is None:
        return ""
    return (user.full_name or user.email or "").strip()


def plant_label(name: str | None, code: str | None) -> str:
    """Plant as shown in reports: "Name (Code)", or whichever part exists."""
    name = (name or "").strip()
    code = (code or "").strip()
    if name and code:
        return f"{name} ({code})"
    return name or code


def _safe_cell(value: str | None) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _clean_pdf_text(value: str) -> str:
    text = _safe_cell(value)
    text = text.replace("\n", " ").replace("\r", " ")
    text = re.sub(r"\s+", " ", text)
    return escape(text)


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "header": ParagraphStyle(
            "ReportHeaderCell",
            fontName="Helvetica-Bold",
            fontSize=PDF_HEADER_FONT,
            leading=9,
            textColor=colors.white,
            alignment=TA_LEFT,
        ),
        "cell": ParagraphStyle(
            "ReportBodyCell",
            fontName="Helvetica",
            fontSize=PDF_BODY_FONT,
            leading=9,
            textColor=colors.HexColor("#111827"),
            alignment=TA_LEFT,
        ),
        "summary_label": ParagraphStyle(
            "ReportSummaryLabel",
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#374151"),
            alignment=TA_LEFT,
        ),
        "summary_value": ParagraphStyle(
            "ReportSummaryValue",
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#111827"),
            alignment=TA_LEFT,
        ),
        "empty": ParagraphStyle(
            "ReportEmpty",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#6B7280"),
        ),
    }


def _column_weights(headers: list[str], rows: list[list[str]]) -> list[float]:
    weights: list[float] = []
    for idx, header in enumerate(headers):
        key = _header_key(header)
        content_cap = 22 if key == "asset name" else 18 if key in _ID_HEADERS else 48
        max_len = len(header)
        for row in rows:
            if idx < len(row):
                cell = _safe_cell(row[idx])
                max_len = max(max_len, min(len(cell), content_cap))
        if key == "asset name":
            max_len = min(max(max_len, 16), 22)
        if key in _ID_HEADERS:
            max_len = max(max_len, 14)
        if key in _FLEXIBLE_HEADERS and key != "asset name":
            max_len = max(max_len, 24)
        if key in _COMPACT_HEADERS:
            max_len = min(max_len, 10)
        if key in {"last verification", "scan time", "created at", "completed at"}:
            max_len = max(max_len, 14)
        weights.append(float(max(max_len, 6)))
    return weights


def _column_widths(headers: list[str], rows: list[list[str]], available_width: float) -> list[float]:
    weights = _column_weights(headers, rows)
    bounds = [_column_bounds(header, available_width) for header in headers]
    mins = [lo for lo, _ in bounds]
    maxs = [hi for _, hi in bounds]
    total_weight = sum(weights) or 1.0

    widths = [mins[i] + (available_width - sum(mins)) * (weights[i] / total_weight) for i in range(len(headers))]

    for _ in range(12):
        for i in range(len(widths)):
            widths[i] = max(mins[i], min(maxs[i], widths[i]))
        delta = available_width - sum(widths)
        if abs(delta) < 0.5:
            break
        growable = [
            (i, weights[i])
            for i in range(len(widths))
            if delta > 0 and widths[i] < maxs[i] - 0.5
        ]
        if not growable:
            if delta < 0:
                scale = available_width / sum(widths)
                widths = [w * scale for w in widths]
            break
        grow_weight = sum(w for _, w in growable)
        for i, weight in growable:
            share = delta * (weight / grow_weight)
            widths[i] = min(maxs[i], widths[i] + share)

    remainder = available_width - sum(widths)
    if abs(remainder) > 0.5:
        sink = [
            i
            for i in range(len(widths))
            if (remainder > 0 and widths[i] < maxs[i] - 0.5) or (remainder < 0 and widths[i] > mins[i] + 0.5)
        ]
        if not sink:
            sink = list(range(len(widths)))
        share = remainder / len(sink)
        for i in sink:
            widths[i] += share

    return widths


def _pdf_data_table(
    document: ReportDocument,
    styles: dict[str, ParagraphStyle],
    available_width: float,
) -> Table:
    headers = ["Sr. No.", *document.headers]
    rows = [[str(idx), *row] for idx, row in enumerate(document.rows, start=1)]
    col_widths = _column_widths(headers, rows, available_width)
    header_row = [Paragraph(_clean_pdf_text(h), styles["header"]) for h in headers]
    body_rows = [
        [Paragraph(_clean_pdf_text(cell), styles["cell"]) for cell in row]
        for row in rows
    ]
    table = Table([header_row, *body_rows], colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), BRAND_RED),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ROW_ALT]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def build_xlsx(document: ReportDocument) -> bytes:
    from modules.platform.settings.document_format import get_document_format

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    headers = ["Sr. No.", *document.headers]
    ncols = max(len(headers), 4)
    write_xlsx_header(
        ws,
        title=document.title,
        subtitle=document.subtitle,
        ncols=ncols,
        generated_by=document.generated_by,
    )

    # Summary stays on PDF only — Excel heading already has generated by/on.
    ws.append(xlsx_pad_row([], ncols=ncols))
    ws.append(xlsx_pad_row(headers, ncols=ncols))
    header_row = ws.max_row

    for idx, row in enumerate(document.rows, start=1):
        cells = [idx, *(_safe_cell(v) for v in row)]
        ws.append(xlsx_pad_row(cells, ncols=ncols))

    data_last_row = ws.max_row

    show_sigs = (
        document.show_signatures
        if document.show_signatures is not None
        else get_document_format().signatures.reports
    )
    if show_sigs:
        _xlsx_signature_rows(ws, document.signature_blocks, ncols=ncols)

    apply_xlsx_data_layout(
        ws, ncols=ncols, header_row=header_row, data_last_row=data_last_row
    )
    apply_xlsx_print_chrome(ws, ncols=ncols, generated_by=document.generated_by)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_signature_rows(ws, blocks: tuple[tuple[str, str, str], ...], *, ncols: int) -> None:
    items = blocks or (("Prepared by", "", ""), ("Approved by", "", ""))
    start = XLSX_LEFT_PAD_COLS + 1
    ws.append(xlsx_pad_row([], ncols=ncols))
    ws.append(xlsx_pad_row([label for label, _n, _d in items], ncols=ncols))
    for col_idx in range(start, start + len(items)):
        ws.cell(row=ws.max_row, column=col_idx).font = Font(bold=True, size=9)
    ws.append(xlsx_pad_row([name or "" for _l, name, _d in items], ncols=ncols))
    for col_idx in range(start, start + len(items)):
        ws.cell(row=ws.max_row, column=col_idx).font = Font(italic=True, size=11)
    ws.append(xlsx_pad_row([f"Name: {name}" if name else "Name:" for _l, name, _d in items], ncols=ncols))
    ws.append(xlsx_pad_row([f"Date: {date}" if date else "Date:" for _l, _n, date in items], ncols=ncols))


def _pdf_summary_block(
    document: ReportDocument,
    styles: dict[str, ParagraphStyle],
    available_width: float,
) -> Table | None:
    if not document.summary:
        return None
    rows = [
        [
            Paragraph(_clean_pdf_text(label), styles["summary_label"]),
            Paragraph(_clean_pdf_text(value), styles["summary_value"]),
        ]
        for label, value in document.summary
    ]
    table = Table(rows, colWidths=[available_width * 0.32, available_width * 0.68])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F4F6")),
                ("BACKGROUND", (1, 0), (1, -1), colors.white),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#374151")),
                ("TEXTCOLOR", (1, 0), (1, -1), colors.HexColor("#111827")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def build_pdf(document: ReportDocument) -> bytes:
    from modules.platform.settings.document_format import get_document_format

    buf = io.BytesIO()
    page_size = A4 if document.force_portrait or len(document.headers) <= 4 else landscape(A4)
    generated_at = format_generated_at()
    show_sigs = (
        document.show_signatures
        if document.show_signatures is not None
        else get_document_format().signatures.reports
    )
    bottom = PDF_MARGIN + FOOTER_BAND + (SIG_BAND if show_sigs else 0)
    pdf = SimpleDocTemplate(
        buf,
        pagesize=page_size,
        leftMargin=PDF_MARGIN,
        rightMargin=PDF_MARGIN,
        topMargin=PDF_MARGIN + HEADER_BAND,
        bottomMargin=bottom,
        title=document.title,
        author=document.generated_by,
        subject=document.subtitle,
    )
    styles = _pdf_styles()
    available_width = page_size[0] - pdf.leftMargin - pdf.rightMargin

    story: list = []
    summary_table = _pdf_summary_block(document, styles, available_width)
    if summary_table is not None:
        story.extend([summary_table, Spacer(1, 12)])

    if not document.rows:
        story.append(Paragraph("No records found for the selected filters.", styles["empty"]))
    else:
        story.append(_pdf_data_table(document, styles, available_width))

    def canvasmaker(*args, **kwargs):
        return ChromeCanvas(
            *args,
            header_title=document.title,
            header_subtitle=document.subtitle,
            generated_at=generated_at,
            generated_by=document.generated_by,
            signature_blocks=document.signature_blocks,
            show_signatures=show_sigs,
            **kwargs,
        )

    pdf.build(story, canvasmaker=canvasmaker)
    return buf.getvalue()


def _filename_token(value: str, *, fallback: str = "Report") -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "", (value or "").strip())
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ._")
    return cleaned or fallback


def report_filename(
    report_key: str,
    fmt: str,
    *,
    report_name: str | None = None,
    financial_year: str | None = None,
) -> str:
    """{Report name}_{DD-MM-YYYY}.ext"""
    _ = financial_year
    stamp = app_now().strftime("%d-%m-%Y")
    ext = "xlsx" if fmt == "xlsx" else "pdf"
    name = _filename_token(report_name or report_key.replace("_", " ").title())
    return f"{name}_{stamp}.{ext}"


def content_type_for_format(fmt: str) -> str:
    if fmt == "xlsx":
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return "application/pdf"
