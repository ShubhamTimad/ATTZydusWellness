"""Generate Excel import template with master-data dropdowns and dependent lists."""

from __future__ import annotations

import io
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from modules.asset_masters.model import (
    ActualLocationMaster,
    AssetClassMaster,
    AssetConditionMaster,
    CostCenterMaster,
    FunctionalLocationMaster,
    ProfitCenterMaster,
    TagColorMaster,
)
from modules.assets.field_definitions import ASSET_IMPORT_TEMPLATE_FIELDS, import_field_header_label
from modules.assets.import_master_resolve import master_import_label
from modules.assets.model import AssetCategory
from modules.platform.org_units.model import OrgUnit
from modules.platform.users.model import User

IMPORT_SHEET = "Import"
LISTS_SHEET = "_Lists"
INSTRUCTIONS_SHEET = "Instructions"
DATA_ROW_START = 2
DATA_ROW_COUNT = 500
DATA_ROW_END = DATA_ROW_START + DATA_ROW_COUNT - 1


def _safe_code(code: str) -> str:
    s = re.sub(r"[^A-Za-z0-9_]", "_", (code or "").strip())
    if not s:
        return "EMPTY"
    if s[0].isdigit():
        s = f"_{s}"
    return s


def _unique_range_name(base: str, used: set[str]) -> str:
    name = base[:255]
    if name not in used:
        used.add(name)
        return name
    n = 2
    while True:
        suffix = f"_{n}"
        candidate = f"{base[: 255 - len(suffix)]}{suffix}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        n += 1


def _excel_sanitize_expr(cell_ref: str) -> str:
    expr = cell_ref
    for ch in ("-", " ", "."):
        expr = f'SUBSTITUTE({expr},"{ch}","_")'
    return expr


def _write_vertical_list(ws, col: int, values: list[str], start_row: int = 2) -> tuple[int, int]:
    for i, value in enumerate(values):
        ws.cell(row=start_row + i, column=col, value=value)
    if not values:
        return start_row, start_row
    return start_row, start_row + len(values) - 1


def _define_list_range(wb: Workbook, name: str, sheet: str, col: int, row_start: int, row_end: int) -> None:
    col_letter = get_column_letter(col)
    attr = f"'{sheet}'!${col_letter}${row_start}:${col_letter}${row_end}"
    wb.defined_names.add(DefinedName(name, attr_text=attr))


def _add_list_validation(
    ws,
    *,
    col: int,
    formula: str,
    allow_blank: bool = True,
) -> None:
    col_letter = get_column_letter(col)
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
    dv.error = "Choose a value from the list."
    dv.errorTitle = "Invalid value"
    dv.prompt = "Select from the dropdown."
    dv.promptTitle = "Allowed values"
    dv.add(f"{col_letter}{DATA_ROW_START}:{col_letter}{DATA_ROW_END}")
    ws.add_data_validation(dv)


def _fill_pair_lookup_formulas(
    ws,
    *,
    source_col: int,
    target_col: int,
    source_range: str,
    target_range: str,
) -> None:
    """Fill target from source via INDEX/MATCH for data rows (e.g. code → name)."""
    source_letter = get_column_letter(source_col)
    for row in range(DATA_ROW_START, DATA_ROW_END + 1):
        src = f"{source_letter}{row}"
        ws.cell(
            row=row,
            column=target_col,
            value=(
                f'=IF({src}="","",IFERROR(INDEX({target_range},MATCH({src},{source_range},0)),""))'
            ),
        )


def _fill_plant_scoped_cc_name_formulas(
    ws,
    *,
    plant_col: int,
    cc_code_col: int,
    cc_name_col: int,
) -> None:
    """Cost Center Name = INDEX(CCN_<plant>, MATCH(code, CC_<plant>))."""
    plant_letter = get_column_letter(plant_col)
    code_letter = get_column_letter(cc_code_col)
    for row in range(DATA_ROW_START, DATA_ROW_END + 1):
        plant_ref = f"{plant_letter}{row}"
        code_ref = f"{code_letter}{row}"
        safe_plant = _excel_sanitize_expr(plant_ref)
        ws.cell(
            row=row,
            column=cc_name_col,
            value=(
                f'=IF(OR({plant_ref}="",{code_ref}=""),"",'
                f'IFERROR(INDEX(INDIRECT("CCN_"&{safe_plant}),'
                f'MATCH({code_ref},INDIRECT("CC_"&{safe_plant}),0)),""))'
            ),
        )


def build_import_template_xlsx(db: Session) -> bytes:
    wb = Workbook()
    ws_import = wb.active
    ws_import.title = IMPORT_SHEET
    ws_lists = wb.create_sheet(LISTS_SHEET)
    ws_instructions = wb.create_sheet(INSTRUCTIONS_SHEET)

    from modules.assets.field_config import AssetFieldConfigService

    svc = AssetFieldConfigService(db)
    fields = [key for key in ASSET_IMPORT_TEMPLATE_FIELDS if svc.is_visible(key)]
    headers = [import_field_header_label(key, db=db) for key in fields]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_import.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(col_idx)
        ws_import.column_dimensions[letter].width = max(12, min(len(header) + 2, 28))

    col_map = {key: idx for idx, key in enumerate(fields, start=1)}
    used_names: set[str] = set()
    list_col = 1

    plants = list(db.scalars(select(OrgUnit).where(OrgUnit.type == "PLANT").order_by(OrgUnit.code)).all())
    # Keep code/name lists aligned for INDEX/MATCH autofill formulas.
    plant_codes = [str(p.code) for p in plants]
    plant_names = [(str(p.name).strip() or str(p.code)) for p in plants]
    row_start, row_end = _write_vertical_list(ws_lists, list_col, plant_codes)
    if plant_codes:
        _define_list_range(wb, "PlantCodes", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1
    row_start, row_end = _write_vertical_list(ws_lists, list_col, plant_names)
    if plant_names:
        _define_list_range(wb, "PlantNames", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1

    categories = list(
        db.scalars(
            select(AssetCategory)
            .where(AssetCategory.is_active.is_(True))
            .order_by(AssetCategory.code)
        ).all()
    )
    category_labels = [
        master_import_label(name=str(c.name), code=str(c.code)) for c in categories
    ]
    row_start, row_end = _write_vertical_list(ws_lists, list_col, category_labels)
    if category_labels:
        _define_list_range(wb, "Categories", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1

    tag_colors = list(
        db.scalars(
            select(TagColorMaster)
            .where(TagColorMaster.is_active.is_(True))
            .order_by(TagColorMaster.code)
        ).all()
    )
    tag_color_labels = [
        master_import_label(name=str(t.name), code=str(t.code)) for t in tag_colors
    ]
    row_start, row_end = _write_vertical_list(ws_lists, list_col, tag_color_labels)
    if tag_color_labels:
        _define_list_range(wb, "TagColors", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1

    custodians = list(
        db.scalars(
            select(User)
            .where(User.is_active.is_(True))
            .order_by(User.full_name, User.username)
        ).all()
    )
    custodian_labels = [
        f"{c.username} — {c.full_name}" for c in custodians
    ]
    row_start, row_end = _write_vertical_list(ws_lists, list_col, custodian_labels)
    if custodian_labels:
        _define_list_range(wb, "AssetCustodians", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1

    conditions = list(
        db.scalars(
            select(AssetConditionMaster)
            .where(AssetConditionMaster.is_active.is_(True))
            .order_by(AssetConditionMaster.code)
        ).all()
    )
    condition_labels = [
        master_import_label(name=str(c.name), code=str(c.code)) for c in conditions
    ]
    row_start, row_end = _write_vertical_list(ws_lists, list_col, condition_labels)
    if condition_labels:
        _define_list_range(wb, "AssetConditions", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1

    asset_classes = list(
        db.scalars(
            select(AssetClassMaster)
            .where(AssetClassMaster.is_active.is_(True))
            .order_by(AssetClassMaster.code)
        ).all()
    )
    asset_class_labels = [
        master_import_label(name=str(c.name), code=str(c.code)) for c in asset_classes
    ]
    row_start, row_end = _write_vertical_list(ws_lists, list_col, asset_class_labels)
    if asset_class_labels:
        _define_list_range(wb, "AssetClasses", LISTS_SHEET, list_col, row_start, row_end)
    list_col += 1

    cost_centers = list(
        db.scalars(
            select(CostCenterMaster)
            .where(CostCenterMaster.is_active.is_(True))
            .order_by(CostCenterMaster.plant_id, CostCenterMaster.code)
        ).all()
    )
    cc_by_plant: dict[int, list[CostCenterMaster]] = defaultdict(list)
    for cc in cost_centers:
        cc_by_plant[int(cc.plant_id)].append(cc)

    plant_by_id = {int(p.id): p for p in plants}
    # Cost Center Code dropdowns (codes); Profit Center lists key off the selected code.
    for plant in plants:
        labels = [str(cc.code) for cc in cc_by_plant.get(int(plant.id), [])]
        col = list_col
        row_start, row_end = _write_vertical_list(ws_lists, col, labels)
        if labels:
            range_name = _unique_range_name(
                f"CC_{_safe_code(str(plant.code))}",
                used_names,
            )
            _define_list_range(wb, range_name, LISTS_SHEET, col, row_start, row_end)
        list_col += 1

    # Cost Center Name dropdowns (display names; import resolves name → stored code).
    for plant in plants:
        labels = [
            master_import_label(name=str(cc.name), code=str(cc.code))
            for cc in cc_by_plant.get(int(plant.id), [])
        ]
        col = list_col
        row_start, row_end = _write_vertical_list(ws_lists, col, labels)
        if labels:
            range_name = _unique_range_name(
                f"CCN_{_safe_code(str(plant.code))}",
                used_names,
            )
            _define_list_range(wb, range_name, LISTS_SHEET, col, row_start, row_end)
        list_col += 1

    profit_centers = list(
        db.scalars(
            select(ProfitCenterMaster)
            .where(ProfitCenterMaster.is_active.is_(True))
            .order_by(ProfitCenterMaster.cost_center_id, ProfitCenterMaster.code)
        ).all()
    )
    pc_by_cc: dict[int, list[ProfitCenterMaster]] = defaultdict(list)
    for pc in profit_centers:
        pc_by_cc[int(pc.cost_center_id)].append(pc)

    for cc in cost_centers:
        plant = plant_by_id.get(int(cc.plant_id))
        if plant is None:
            continue
        labels = [
            master_import_label(name=str(pc.name), code=str(pc.code))
            for pc in pc_by_cc.get(int(cc.id), [])
        ]
        col = list_col
        row_start, row_end = _write_vertical_list(ws_lists, col, labels)
        if labels:
            range_name = _unique_range_name(
                f"PC_{_safe_code(str(plant.code))}__{_safe_code(str(cc.code))}",
                used_names,
            )
            _define_list_range(wb, range_name, LISTS_SHEET, col, row_start, row_end)
        list_col += 1

    functional_locations = list(
        db.scalars(
            select(FunctionalLocationMaster)
            .where(FunctionalLocationMaster.is_active.is_(True))
            .order_by(FunctionalLocationMaster.plant_id, FunctionalLocationMaster.code)
        ).all()
    )
    fl_by_plant: dict[int, list[FunctionalLocationMaster]] = defaultdict(list)
    for fl in functional_locations:
        fl_by_plant[int(fl.plant_id)].append(fl)

    for plant in plants:
        labels = [
            master_import_label(name=str(fl.name), code=str(fl.code))
            for fl in fl_by_plant.get(int(plant.id), [])
        ]
        col = list_col
        row_start, row_end = _write_vertical_list(ws_lists, col, labels)
        if labels:
            range_name = _unique_range_name(
                f"FL_{_safe_code(str(plant.code))}",
                used_names,
            )
            _define_list_range(wb, range_name, LISTS_SHEET, col, row_start, row_end)
        list_col += 1

    actual_locations = list(
        db.scalars(
            select(ActualLocationMaster)
            .where(ActualLocationMaster.is_active.is_(True))
            .order_by(ActualLocationMaster.functional_location_id, ActualLocationMaster.code)
        ).all()
    )
    al_by_fl: dict[int, list[ActualLocationMaster]] = defaultdict(list)
    for al in actual_locations:
        al_by_fl[int(al.functional_location_id)].append(al)

    fl_by_id = {int(fl.id): fl for fl in functional_locations}
    for fl in functional_locations:
        plant = plant_by_id.get(int(fl.plant_id))
        if plant is None:
            continue
        labels = [
            master_import_label(name=str(al.name), code=str(al.code))
            for al in al_by_fl.get(int(fl.id), [])
        ]
        col = list_col
        row_start, row_end = _write_vertical_list(ws_lists, col, labels)
        if labels:
            range_name = _unique_range_name(
                f"AL_{_safe_code(str(plant.code))}__{_safe_code(master_import_label(name=str(fl.name), code=str(fl.code)))}",
                used_names,
            )
            _define_list_range(wb, range_name, LISTS_SHEET, col, row_start, row_end)
        list_col += 1

    ws_lists.sheet_state = "hidden"

    if "plant_code" in col_map and plant_codes:
        _add_list_validation(ws_import, col=col_map["plant_code"], formula="=PlantCodes")

    if "plant_name" in col_map and plant_names:
        _add_list_validation(ws_import, col=col_map["plant_name"], formula="=PlantNames")

    if (
        "plant_code" in col_map
        and "plant_name" in col_map
        and plant_codes
        and plant_names
    ):
        _fill_pair_lookup_formulas(
            ws_import,
            source_col=col_map["plant_code"],
            target_col=col_map["plant_name"],
            source_range="PlantCodes",
            target_range="PlantNames",
        )

    if "category_code" in col_map and category_labels:
        _add_list_validation(ws_import, col=col_map["category_code"], formula="=Categories")

    if "tag_color_code" in col_map and tag_color_labels:
        _add_list_validation(ws_import, col=col_map["tag_color_code"], formula="=TagColors")

    if "asset_custodian_code" in col_map and custodian_labels:
        _add_list_validation(ws_import, col=col_map["asset_custodian_code"], formula="=AssetCustodians")

    if "asset_condition_code" in col_map and condition_labels:
        _add_list_validation(ws_import, col=col_map["asset_condition_code"], formula="=AssetConditions")
    if "asset_class_code" in col_map and asset_class_labels:
        _add_list_validation(ws_import, col=col_map["asset_class_code"], formula="=AssetClasses")

    if "plant_code" in col_map and "cost_center" in col_map and plant_codes:
        plant_ref = f"{get_column_letter(col_map['plant_code'])}2"
        formula = (
            f'=IF({plant_ref}="","",INDIRECT("CC_"&{_excel_sanitize_expr(plant_ref)}))'
        )
        _add_list_validation(ws_import, col=col_map["cost_center"], formula=formula)

    if "plant_code" in col_map and "cost_center_name" in col_map and plant_codes:
        plant_ref = f"{get_column_letter(col_map['plant_code'])}2"
        formula = (
            f'=IF({plant_ref}="","",INDIRECT("CCN_"&{_excel_sanitize_expr(plant_ref)}))'
        )
        _add_list_validation(ws_import, col=col_map["cost_center_name"], formula=formula)

    if all(k in col_map for k in ("plant_code", "cost_center", "cost_center_name")) and plant_codes:
        _fill_plant_scoped_cc_name_formulas(
            ws_import,
            plant_col=col_map["plant_code"],
            cc_code_col=col_map["cost_center"],
            cc_name_col=col_map["cost_center_name"],
        )

    if all(k in col_map for k in ("plant_code", "cost_center", "profit_center")):
        plant_ref = f"{get_column_letter(col_map['plant_code'])}2"
        cc_ref = f"{get_column_letter(col_map['cost_center'])}2"
        formula = (
            f'=IF(OR({plant_ref}="",{cc_ref}=""),"",'
            f'INDIRECT("PC_"&{_excel_sanitize_expr(plant_ref)}&"__"&{_excel_sanitize_expr(cc_ref)}))'
        )
        _add_list_validation(ws_import, col=col_map["profit_center"], formula=formula)

    if "plant_code" in col_map and "functional_location" in col_map and plant_codes:
        plant_ref = f"{get_column_letter(col_map['plant_code'])}2"
        formula = (
            f'=IF({plant_ref}="","",INDIRECT("FL_"&{_excel_sanitize_expr(plant_ref)}))'
        )
        _add_list_validation(ws_import, col=col_map["functional_location"], formula=formula)

    if all(k in col_map for k in ("plant_code", "functional_location", "actual_location")):
        plant_ref = f"{get_column_letter(col_map['plant_code'])}2"
        fl_ref = f"{get_column_letter(col_map['functional_location'])}2"
        formula = (
            f'=IF(OR({plant_ref}="",{fl_ref}=""),"",'
            f'INDIRECT("AL_"&{_excel_sanitize_expr(plant_ref)}&"__"&{_excel_sanitize_expr(fl_ref)}))'
        )
        _add_list_validation(ws_import, col=col_map["actual_location"], formula=formula)

    ws_import.freeze_panes = "A2"

    instructions = [
        ("Asset import template", ""),
        ("", ""),
        ("How to use", ""),
        ("1.", "Fill one row per asset on the Import sheet. Do not change header row names."),
        (
            "2.",
            "Required (*): Asset name, Combined Asset No, Asset category, Plant Code (or Plant Name), "
            "and Cost Center Code (or Cost Center Name). Actual Location 1 / 2 are optional. "
            "RFID is optional but must be unique when set.",
        ),
        (
            "3.",
            "Plant Code ↔ Plant Name: select Plant Code and Plant Name fills automatically. "
            "Or select Plant Name alone (code is resolved on upload). "
            "Select Plant Code first so Cost center / Profit center / Actual Location 1 lists filter correctly.",
        ),
        (
            "4.",
            "Cost Center Code ↔ Cost Center Name: select Cost Center Code and Name fills automatically. "
            "Or select Name alone (code is resolved on upload). "
            "Select Cost Center Code before Profit center. Select Actual Location 1 before Actual location 2.",
        ),
        (
            "5.",
            "CapEx form number is optional; use an approved, unused form (e.g. CAP-2026-0001). "
            "Tag color is only valid when RFID is provided.",
        ),
        ("6.", "Dates: use YYYY-MM-DD (e.g. 2024-06-15)."),
        ("7.", "Save as .xlsx and upload via Assets → Import."),
        ("", ""),
        ("Master data in this file", f"{len(plant_codes)} plants · {len(category_labels)} categories · "
         f"{len(cost_centers)} cost centers · {len(profit_centers)} profit centers · "
         f"{len(functional_locations)} functional locations · {len(actual_locations)} actual locations"),
    ]
    for row_idx, (left, right) in enumerate(instructions, start=1):
        ws_instructions.cell(row=row_idx, column=1, value=left).font = Font(bold=row_idx == 1)
        ws_instructions.cell(row=row_idx, column=2, value=right)
    ws_instructions.column_dimensions["A"].width = 28
    ws_instructions.column_dimensions["B"].width = 72

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
