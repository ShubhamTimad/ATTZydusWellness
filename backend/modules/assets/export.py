"""Excel export for asset register."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook

from modules.assets.field_definitions import ASSET_EXPORT_FIELD_LABELS, ASSET_REGISTER_PDF_LABELS
from modules.assets.model import AssetMaster
from modules.reports.pdf_chrome import (
    apply_xlsx_data_layout,
    apply_xlsx_print_chrome,
    write_xlsx_header,
    xlsx_pad_row,
)


def _format_cell(value: Any) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def asset_export_value(asset: AssetMaster, field: str) -> Any:
    loc = asset.location
    fin = asset.financial
    inv = asset.inventory
    plant = loc.plant if loc else None
    category = asset.category
    tag = asset.tag_color
    custodian = asset.asset_custodian
    condition = asset.asset_condition
    asset_class = asset.asset_class_master

    values: dict[str, Any] = {
        "asset_name": asset.asset_name,
        "asset_class_code": (asset_class.name or asset_class.code) if asset_class else asset.asset_class,
        "asset_type": asset.asset_type,
        "asset_code": asset.asset_code,
        "combined_asset_number": asset.combined_asset_number,
        "remark": asset.remark,
        "sub_no": asset.sub_no,
        "rfid_code": asset.rfid_code,
        "serial_no": asset.serial_no,
        "asset_description": asset.information.asset_description if asset.information else None,
        "category_code": category.code if category else None,
        "category_name": category.name if category else None,
        "tag_color_code": tag.code if tag else None,
        "asset_custodian_code": custodian.username if custodian else None,
        "asset_condition_code": condition.code if condition else None,
        "sub_location": loc.sub_location if loc else None,
        "functional_location": loc.functional_location if loc else None,
        "actual_location": loc.actual_location if loc else None,
        "plant_code": plant.code if plant else None,
        "plant_name": plant.name if plant else None,
        "cost_center": fin.cost_center if fin else None,
        "party_details": fin.party_details if fin else None,
        "invoice_no": fin.invoice_no if fin else None,
        "invoice_date": fin.invoice_date if fin else None,
        "voucher_no": fin.voucher_no if fin else None,
        "voucher_date": fin.voucher_date if fin else None,
        "grn_date": fin.grn_date if fin else None,
        "grn_no": fin.grn_no if fin else None,
        "put_to_use_date": fin.put_to_use_date if fin else None,
        "depreciation_rate": fin.depreciation_rate if fin else None,
        "gross_block": fin.gross_block if fin else None,
        # ponytail: no stored current APC; blank until a source field exists.
        "current_apc_igaap": None,
        "life_span": fin.life_span if fin else None,
        "profit_center": fin.profit_center if fin else None,
        "inventory_count": inv.inventory_count if inv else 1,
        "physical_count": inv.physical_count if inv else None,
        "last_verification_date": inv.last_verification_date if inv else None,
        "cap_date": fin.cap_date if fin else None,
        "pm_equipment_id": asset.information.pm_equipment_id if asset.information else None,
        "cost_center_name": None,
        "capitalization_form_number": None,
        "fy_verification_count": None,
        "physical_order_id": None,
        "physical_order_status": None,
        "lifecycle_status": asset.lifecycle_status,
        "approval_status": asset.approval_status,
        "source_type": asset.source_type,
    }
    return values.get(field)


def build_asset_export_xlsx(
    assets: list[AssetMaster], fields: list[str], *, generated_by: str = ""
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"

    headers = ["Sr. No."] + [
        ASSET_EXPORT_FIELD_LABELS.get(key, key.replace("_", " ").title()) for key in fields
    ]
    ncols = max(len(headers), 4)
    count = len(assets)
    write_xlsx_header(
        ws,
        title="Asset Register",
        subtitle=f"{count} asset{'s' if count != 1 else ''}",
        ncols=ncols,
        generated_by=generated_by,
    )
    ws.append(xlsx_pad_row([], ncols=ncols))
    ws.append(xlsx_pad_row(headers, ncols=ncols))
    header_row = ws.max_row

    for idx, asset in enumerate(assets, start=1):
        row = [idx] + [_format_cell(asset_export_value(asset, key)) for key in fields]
        ws.append(xlsx_pad_row(row, ncols=ncols))

    apply_xlsx_data_layout(ws, ncols=ncols, header_row=header_row)
    apply_xlsx_print_chrome(ws, ncols=ncols, generated_by=generated_by)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_asset_export_pdf(
    assets: list[AssetMaster], fields: list[str], *, generated_by: str = ""
) -> bytes:
    from modules.reports.export import ReportDocument, build_pdf

    headers = [
        ASSET_REGISTER_PDF_LABELS.get(key) or ASSET_EXPORT_FIELD_LABELS.get(key, key.replace("_", " ").title())
        for key in fields
    ]
    rows: list[list[str]] = []
    for asset in assets:
        row: list[str] = []
        for key in fields:
            cell = _format_cell(asset_export_value(asset, key))
            row.append("" if cell is None else str(cell))
        rows.append(row)
    return build_pdf(
        ReportDocument(
            title="Asset Register",
            subtitle=f"{len(rows)} asset{'s' if len(rows) != 1 else ''}",
            headers=headers,
            rows=rows,
            generated_by=generated_by,
            force_portrait=True,
        )
    )
