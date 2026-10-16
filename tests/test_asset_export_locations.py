"""Asset register export shows location names, not stored codes."""

from __future__ import annotations

import base64
import io
import re
import zlib
from datetime import date
from types import SimpleNamespace

from openpyxl import load_workbook

from modules.assets.export import asset_export_value, build_asset_export_pdf, build_asset_export_xlsx
from modules.platform.settings.document_format import DEFAULTS, previewing


def _asset(*, plant_id=1, functional_location="1", actual_location="2"):
    loc = SimpleNamespace(
        plant_id=plant_id,
        functional_location=functional_location,
        actual_location=actual_location,
        sub_location=None,
        plant=SimpleNamespace(code="ZT001", name="ZYDUS_TAKEDA"),
    )
    return SimpleNamespace(
        location=loc,
        financial=None,
        inventory=None,
        category=None,
        tag_color=None,
        asset_custodian=None,
        asset_condition=None,
        asset_class_master=None,
        information=None,
        asset_name="Chair",
        asset_class=None,
        asset_type=None,
        asset_code="A1",
        combined_asset_number=None,
        remark=None,
        sub_no=None,
        rfid_code=None,
        serial_no=None,
        lifecycle_status=None,
        approval_status=None,
        source_type=None,
    )


def test_export_uses_actual_location_names() -> None:
    asset = _asset()
    fl_desc = {(1, "1"): "Production Floor"}
    al_desc = {(1, "1", "2"): "Reactor Bay"}
    assert (
        asset_export_value(asset, "functional_location", fl_desc=fl_desc, al_desc=al_desc)
        == "Production Floor"
    )
    assert (
        asset_export_value(asset, "actual_location", fl_desc=fl_desc, al_desc=al_desc)
        == "Reactor Bay"
    )


def test_export_falls_back_to_code_when_name_missing() -> None:
    asset = _asset()
    assert asset_export_value(asset, "functional_location") == "1"
    assert asset_export_value(asset, "actual_location") == "2"


def test_export_uses_category_and_condition_names() -> None:
    asset = _asset()
    asset.category = SimpleNamespace(code="CAT", name="IT Equipment")
    asset.asset_condition = SimpleNamespace(code="GD", name="Good")
    asset.asset_custodian = SimpleNamespace(username="jdoe", full_name="Jane Doe")
    assert asset_export_value(asset, "category_code") == "IT Equipment"
    assert asset_export_value(asset, "asset_condition_code") == "Good"
    assert asset_export_value(asset, "asset_custodian_code") == "Jane Doe"


def test_export_master_name_falls_back_to_code() -> None:
    asset = _asset()
    asset.category = SimpleNamespace(code="CAT", name="")
    asset.asset_condition = SimpleNamespace(code="GD", name=None)
    assert asset_export_value(asset, "category_code") == "CAT"
    assert asset_export_value(asset, "asset_condition_code") == "GD"


def test_export_capex_number_and_app_date_format(monkeypatch) -> None:
    import modules.reports.pdf_chrome as chrome

    monkeypatch.setattr(chrome, "_chrome_strings", lambda: ("Takeda", "Confidential", "Takeda"))
    monkeypatch.setattr(chrome, "logo_png_bytes", lambda: None)
    asset = _asset()
    asset.id = 7
    asset.financial = SimpleNamespace(
        cap_date=date(2026, 9, 1),
        **{k: None for k in (
            "cost_center", "party_details", "invoice_no", "invoice_date", "voucher_no", "voucher_date",
            "grn_date", "grn_no", "put_to_use_date", "depreciation_rate", "gross_block", "life_span",
            "profit_center",
        )},
    )
    asset.inventory = SimpleNamespace(
        last_verification_date=date(2026, 7, 28), inventory_count=1, physical_count=None
    )
    raw = build_asset_export_xlsx(
        [asset],
        ["capitalization_form_number", "cap_date", "last_verification_date"],
        capex_no={7: "CAP-2026-0001"},
    )
    rows = list(load_workbook(io.BytesIO(raw)).active.iter_rows(values_only=True))
    assert any(row and "CAP-2026-0001" in row and "01/09/26" in row and "28/07/26" in row for row in rows)


def _named_asset() -> SimpleNamespace:
    asset = _asset()
    asset.asset_name = "Chair"
    asset.category = SimpleNamespace(code="CAT", name="IT Equipment")
    asset.asset_condition = SimpleNamespace(code="GD", name="Good")
    return asset


def _pdf_text(data: bytes) -> str:
    parts: list[str] = []
    for match in re.finditer(rb"/Length (\d+)\s*>>\s*stream\n", data):
        chunk = data[match.end() : match.end() + int(match.group(1))]
        try:
            chunk = zlib.decompress(base64.a85decode(chunk, adobe=True))
        except (ValueError, zlib.error):
            try:
                chunk = zlib.decompress(chunk)
            except zlib.error:
                pass
        parts.append(chunk.decode("latin1", errors="ignore"))
    return "\n".join(parts)


def test_register_excel_includes_asset_rows(monkeypatch) -> None:
    import modules.reports.pdf_chrome as chrome

    monkeypatch.setattr(chrome, "_chrome_strings", lambda: ("Takeda", "Confidential", "Takeda"))
    monkeypatch.setattr(chrome, "logo_png_bytes", lambda: None)
    raw = build_asset_export_xlsx(
        [_named_asset()],
        ["asset_name", "category_code", "asset_condition_code"],
        generated_by="Tester",
    )
    ws = load_workbook(io.BytesIO(raw)).active
    rows = list(ws.iter_rows(values_only=True))
    assert any(row and "Generated On:" in "".join(str(cell) for cell in row if cell) for row in rows)
    assert any(row and "Chair" in row and "IT Equipment" in row and "Good" in row for row in rows)


def test_register_pdf_includes_sr_no_and_names(monkeypatch) -> None:
    import modules.reports.pdf_chrome as chrome

    monkeypatch.setattr(chrome, "logo_png_bytes", lambda: None)
    with previewing(DEFAULTS):
        raw = build_asset_export_pdf(
            [_named_asset()],
            ["asset_name", "category_code", "asset_condition_code"],
            generated_by="Tester",
        )
    text = _pdf_text(raw)
    assert "Sr. No." in text
    assert "IT Equipment" in text
    assert "Good" in text
    assert "Generated On:" in text
